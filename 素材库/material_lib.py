#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
创作素材库 · 入库工具
================================================================
把从视频/文案中提取的创作素材（人设原型、情节模块、金句表达、冲突主题）
追加到一个可长期积累的 Excel 多表素材库中。

用法：
  python material_lib.py init                  # 初始化空库（已存在则跳过）
  python material_lib.py add <条目.json>       # 追加一批素材
  python material_lib.py stat                  # 查看库内统计

条目 JSON 结构：
{
  "source": {"title": "视频标题", "type": "游戏动漫影视解说", "note": ""},
  "人设原型": [{"原型名称": "...", "核心特征": "...", "内在矛盾": "...",
                "行为逻辑": "...", "弧光走向": "...", "适用题材": "...", "标签": "..."}],
  "情节模块": [{"模块名称": "...", "结构描述": "...", "情绪效果": "...",
                "可套用场景": "...", "适用题材": "...", "标签": "..."}],
  "金句表达": [{"金句": "...", "表达技巧": "...", "使用情境": "...", "标签": "..."}],
  "冲突主题": [{"冲突主题": "...", "描述": "...", "适用题材": "...", "标签": "..."}]
}
"""
import os
import sys
import json
import time
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

def _resolve_lib():
    """素材库路径：环境变量 → 配置文件 → 脚本同目录（默认）"""
    v = os.environ.get('MATERIAL_LIB')
    if v:
        return v
    cfg_path = os.path.join(os.path.expanduser('~'), '.workbuddy-ai',
                            'douyin-data', 'config.json')
    try:
        cfg = json.load(open(cfg_path, encoding='utf-8'))
        if cfg.get('material_lib'):
            return cfg['material_lib']
    except Exception:
        pass
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), '创作素材库.xlsx')


# 素材库位置
LIB = _resolve_lib()

# 表名 -> (列定义, 编号前缀)
SHEETS = {
    '人设原型': (
        ['编号', '原型名称', '核心特征', '内在矛盾', '行为逻辑', '弧光走向',
         '适用题材', '来源类型', '来源标题', '标签', '入库日期'], 'P'),
    '情节模块': (
        ['编号', '模块名称', '结构描述', '情绪效果', '可套用场景',
         '适用题材', '来源类型', '来源标题', '标签', '入库日期'], 'S'),
    '金句表达': (
        ['编号', '金句', '表达技巧', '使用情境',
         '来源类型', '来源标题', '标签', '入库日期'], 'Q'),
    '冲突主题': (
        ['编号', '冲突主题', '描述', '适用题材',
         '来源类型', '来源标题', '标签', '入库日期'], 'C'),
    '来源索引': (
        ['序号', '视频标题', '类型', '入库条目数', '处理日期', '备注'], ''),
}

WIDTHS = {
    '编号': 8, '原型名称': 18, '核心特征': 34, '内在矛盾': 34, '行为逻辑': 34,
    '弧光走向': 30, '适用题材': 16, '来源类型': 18, '来源标题': 30,
    '标签': 16, '入库日期': 12, '模块名称': 18, '结构描述': 44, '情绪效果': 22,
    '可套用场景': 30, '金句': 44, '表达技巧': 26, '使用情境': 26,
    '冲突主题': 18, '描述': 44, '序号': 8, '类型': 18, '入库条目数': 12,
    '处理日期': 12, '备注': 24,
}
HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')


def log(*a):
    print(*a, flush=True)


def init_lib(path=LIB):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        log('库已存在：' + path)
        return path
    wb = Workbook()
    wb.remove(wb.active)
    for name, (cols, _) in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal='center', vertical='center')
        for i, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 20)
        ws.freeze_panes = 'A2'
    wb.save(path)
    log('已初始化素材库：' + path)
    return path


def _next_id(ws, prefix, id_col=1):
    mx = 0
    for row in ws.iter_rows(min_row=2, min_col=id_col, max_col=id_col, values_only=True):
        v = str(row[0] or '')
        if v.startswith(prefix) and v[len(prefix):].isdigit():
            mx = max(mx, int(v[len(prefix):]))
    return f'{prefix}{mx + 1:03d}'


def _existing_keys(ws, key_col):
    keys = set()
    for row in ws.iter_rows(min_row=2, min_col=key_col, max_col=key_col, values_only=True):
        if row[0]:
            keys.add(str(row[0]).strip())
    return keys


def add_batch(path, data):
    if not os.path.exists(path):
        init_lib(path)
    wb = load_workbook(path)
    src = data.get('source', {})
    today = time.strftime('%Y-%m-%d')
    counts = {}
    skipped = []

    for name, (cols, prefix) in SHEETS.items():
        if name == '来源索引':
            continue
        items = data.get(name) or []
        if not items:
            counts[name] = 0
            continue
        ws = wb[name]
        key_col = 2  # 名称类字段在第 2 列
        exist = _existing_keys(ws, key_col)
        n = 0
        for it in items:
            key = str(it.get(cols[1], '')).strip()
            if key and key in exist:
                skipped.append(f'{name}:{key}')
                continue
            row = []
            for col in cols:
                if col == '编号':
                    row.append(_next_id(ws, prefix))
                elif col == '来源类型':
                    row.append(src.get('type', ''))
                elif col == '来源标题':
                    row.append(src.get('title', ''))
                elif col == '入库日期':
                    row.append(today)
                else:
                    row.append(str(it.get(col, '')))
            ws.append(row)
            if key:
                exist.add(key)
            n += 1
        counts[name] = n

    # 来源索引
    ws = wb['来源索引']
    total = sum(counts.values())
    seq = _next_id(ws, '', id_col=1)
    try:
        seq = f'{int(seq):03d}'
    except ValueError:
        seq = seq
    ws.append([seq, src.get('title', ''), src.get('type', ''),
               total, today, src.get('note', '')])

    wb.save(path)
    log(f'入库完成：共 {total} 条 -> ' +
        '，'.join(f'{k} {v}' for k, v in counts.items() if v))
    if skipped:
        log('跳过重复：' + '、'.join(skipped))
    return total


def stat(path=LIB):
    if not os.path.exists(path):
        log('库不存在：' + path)
        return
    wb = load_workbook(path, read_only=True)
    for name in SHEETS:
        ws = wb[name]
        n = ws.max_row - 1 if ws.max_row else 0
        log(f'  {name}: {n} 条')
    wb.close()


def main():
    if len(sys.argv) < 2:
        log(__doc__)
        return
    cmd = sys.argv[1]
    if cmd == 'init':
        init_lib()
    elif cmd == 'add':
        if len(sys.argv) < 3:
            log('用法：add <条目.json>')
            sys.exit(1)
        with open(sys.argv[2], encoding='utf-8') as fh:
            data = json.load(fh)
        add_batch(LIB, data)
    elif cmd == 'stat':
        stat()
    else:
        log('未知命令：' + cmd)


if __name__ == '__main__':
    main()
